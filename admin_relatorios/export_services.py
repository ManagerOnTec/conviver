import io
import re
import unicodedata
import zipfile
from datetime import datetime
from django.utils import timezone
from django.core.files.storage import default_storage
from django.db.models import Q
import PyPDF2
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from documentos_legais.models import DocumentoLegalInternacao
from .models import Relatorio


# Mapeamento de categorias de prontuário para pastas organizadas
PASTAS_CATEGORIAS = {
    'termos_legais': '01_Termos_Legais',
    'diagnosticos_sae': '02_Diagnosticos_e_SAE',
    'prescricoes_adep': '03_Prescricoes_e_ADEP',
    'evolucoes': '04_Evolucoes_Clinicas',
    'sinais_vitais_balanco': '05_Sinais_Vitais_e_Balanco',
    'plano_cuidados': '06_Plano_de_Cuidados',
    'outros': '07_Outros_Relatorios',
}

TIPO_PARA_CATEGORIA = {
    # Termos
    'termo': 'termos_legais',
    'termos': 'termos_legais',
    'documentos_legais': 'termos_legais',
    'termo_legal': 'termos_legais',
    # Diagnósticos e SAE
    'diagnostico': 'diagnosticos_sae',
    'sae': 'diagnosticos_sae',
    # Prescrição e ADEP
    'prescricao': 'prescricoes_adep',
    'adep': 'prescricoes_adep',
    # Evoluções
    'evolucao': 'evolucoes',
    'psicoterapia': 'evolucoes',
    # Sinais vitais e balanço
    'sinais_vitais': 'sinais_vitais_balanco',
    'sv': 'sinais_vitais_balanco',
    'perdas_ganhos': 'sinais_vitais_balanco',
    # Cuidados
    'plano_cuidados': 'plano_cuidados',
    'cuidados': 'plano_cuidados',
    # Outros
    'ata': 'outros',
    'atas': 'outros',
    'oficio': 'outros',
    'oficios': 'outros',
    'orcamento': 'outros',
    'orcamentos': 'outros',
}

TIPO_LABELS = {
    'evolucao': 'Evolução Clínica',
    'prescricao': 'Prescrição Médica',
    'adep': 'ADEP (Administração)',
    'sinais_vitais': 'Sinais Vitais',
    'sae': 'SAE (Enfermagem)',
    'diagnostico': 'Diagnóstico (CID)',
    'plano_cuidados': 'Plano de Cuidados',
    'perdas_ganhos': 'Perdas e Ganhos',
    'documentos_legais': 'Documentos Legais / Termos',
    'ata': 'Ata',
    'atas': 'Atas',
    'oficio': 'Ofício',
    'oficios': 'Ofícios',
    'orcamento': 'Orçamento',
    'orcamentos': 'Orçamentos',
}


def sanitizar_nome(texto):
    """Remove caracteres especiais e espaços para nomes de arquivos seguros."""
    if not texto:
        return 'sem_identificacao'
    # Normaliza unicode (remove acentos)
    nfkd = unicodedata.normalize('NFKD', str(texto))
    sem_acento = u"".join([c for c in nfkd if not unicodedata.combining(c)])
    # Substitui caracteres não alfanuméricos por underline
    limpo = re.sub(r'[^a-zA-Z0-9_-]', '_', sem_acento)
    limpo = re.sub(r'_+', '_', limpo).strip('_')
    return limpo or 'arquivo'


def obter_bytes_pdf(arquivo_field):
    """
    Lê os bytes de um FileField ou caminho de arquivo via default_storage.
    Verifica se existe, tem tamanho > 0 e inicia com %PDF.
    """
    if not arquivo_field:
        return None

    caminho = getattr(arquivo_field, 'name', str(arquivo_field))
    if not caminho:
        return None

    try:
        if not default_storage.exists(caminho):
            return None
        if default_storage.size(caminho) <= 0:
            return None

        with default_storage.open(caminho, 'rb') as f:
            pdf_bytes = f.read()

        if pdf_bytes and pdf_bytes.startswith(b'%PDF'):
            return pdf_bytes
    except Exception as e:
        print(f"[ExportServices] Erro ao ler PDF {caminho}: {e}")
        return None

    return None


def gerar_zip_dossie_pacientes(relatorios_queryset, user=None):
    """
    Gera um arquivo ZIP com o dossiê completo de prontuário e termos legais
    para todos os pacientes representados no queryset de relatórios.
    
    Inclui:
    - Todos os PDFs ativos de Relatorio (Evolução, Prescrição, ADEP, SAE, etc.)
    - Todos os PDFs ativos de DocumentoLegalInternacao
    - Estrutura de pastas categorizada por módulo clínico
    - Manifesto/Sumário textual com data, paciente e lista de arquivos
    """
    # Identifica os pacientes únicos do queryset
    pessoas_dict = {}
    for r in relatorios_queryset.select_related('pessoa', 'atendimento__pessoa'):
        p = r.pessoa or (r.atendimento.pessoa if r.atendimento else None)
        if p and p.id not in pessoas_dict:
            pessoas_dict[p.id] = p

    if not pessoas_dict:
        return None, None, 0

    multiplos_pacientes = len(pessoas_dict) > 1
    total_geral_arquivos = 0
    buffer = io.BytesIO()

    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    data_formatada_extenso = timezone.now().strftime('%d/%m/%Y às %H:%M:%S')

    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for pessoa_id, paciente in pessoas_dict.items():
            nome_paciente_limpo = sanitizar_nome(paciente.nome)
            cpf_paciente = getattr(paciente, 'cpf', '') or 'SemCPF'
            
            # Prefixo da pasta no ZIP (se múltiplos pacientes, cada um tem sua pasta raiz)
            if multiplos_pacientes:
                pasta_raiz_paciente = f"Paciente_{nome_paciente_limpo}_{cpf_paciente}/"
            else:
                pasta_raiz_paciente = ""

            arquivos_do_paciente = []
            nomes_usados_no_zip = set()

            # 1. Busca todos os relatórios ativos ligados a esta pessoa
            relatorios_ativos = Relatorio.objects.filter(
                Q(pessoa=paciente) | Q(atendimento__pessoa=paciente),
                status='A',
            ).select_related('atendimento', 'us_registro').order_by('tipo', '-dt_registro')

            for rel in relatorios_ativos:
                if not rel.relatorio:
                    continue
                pdf_bytes = obter_bytes_pdf(rel.relatorio)
                if not pdf_bytes:
                    continue

                tipo_lower = (rel.tipo or '').lower()
                categoria_chave = TIPO_PARA_CATEGORIA.get(tipo_lower, 'outros')
                pasta_categoria = PASTAS_CATEGORIAS[categoria_chave]

                tipo_nome = TIPO_LABELS.get(tipo_lower, tipo_lower.capitalize())
                data_reg = rel.dt_registro.strftime('%Y%m%d') if rel.dt_registro else 'semdata'
                nome_base = f"{tipo_nome}_{data_reg}_ID{rel.id}.pdf"
                nome_base = sanitizar_nome(nome_base.replace('.pdf', '')) + '.pdf'

                # Evita colisões de nomes dentro da pasta
                caminho_completo = f"{pasta_raiz_paciente}{pasta_categoria}/{nome_base}"
                contador = 1
                while caminho_completo in nomes_usados_no_zip:
                    nome_alt = f"{nome_base.rsplit('.', 1)[0]}_{contador}.pdf"
                    caminho_completo = f"{pasta_raiz_paciente}{pasta_categoria}/{nome_alt}"
                    contador += 1

                nomes_usados_no_zip.add(caminho_completo)
                zip_file.writestr(caminho_completo, pdf_bytes)
                total_geral_arquivos += 1

                tamanho_kb = round(len(pdf_bytes) / 1024, 1)
                arquivos_do_paciente.append({
                    'categoria': pasta_categoria,
                    'arquivo': caminho_completo.replace(pasta_raiz_paciente, ''),
                    'tipo': tipo_nome,
                    'data': rel.dt_registro.strftime('%d/%m/%Y %H:%M') if rel.dt_registro else '-',
                    'usuario': rel.us_registro.username if rel.us_registro else '-',
                    'tamanho_kb': tamanho_kb,
                })

            # 2. Busca todos os termos legais ativos do paciente
            documentos_legais = DocumentoLegalInternacao.objects.filter(
                Q(pessoa=paciente) | Q(atendimento__pessoa=paciente),
                status='A',
            ).select_related('tipo_documento', 'us_registro').order_by('-dt_registro')

            for doc in documentos_legais:
                if not doc.pdf_gerado:
                    continue
                pdf_bytes = obter_bytes_pdf(doc.pdf_gerado)
                if not pdf_bytes:
                    continue

                pasta_categoria = PASTAS_CATEGORIAS['termos_legais']
                titulo_termo = sanitizar_nome(doc.titulo_documento or doc.tipo_documento.nome if doc.tipo_documento else 'Termo')
                codigo = doc.codigo_documento or str(doc.id)
                nome_base = f"Termo_{titulo_termo}_{codigo}.pdf"

                caminho_completo = f"{pasta_raiz_paciente}{pasta_categoria}/{nome_base}"
                contador = 1
                while caminho_completo in nomes_usados_no_zip:
                    nome_alt = f"Termo_{titulo_termo}_{codigo}_{contador}.pdf"
                    caminho_completo = f"{pasta_raiz_paciente}{pasta_categoria}/{nome_alt}"
                    contador += 1

                nomes_usados_no_zip.add(caminho_completo)
                zip_file.writestr(caminho_completo, pdf_bytes)
                total_geral_arquivos += 1

                tamanho_kb = round(len(pdf_bytes) / 1024, 1)
                arquivos_do_paciente.append({
                    'categoria': pasta_categoria,
                    'arquivo': caminho_completo.replace(pasta_raiz_paciente, ''),
                    'tipo': f"Termo Legal: {doc.titulo_documento}",
                    'data': doc.dt_registro.strftime('%d/%m/%Y %H:%M') if doc.dt_registro else '-',
                    'usuario': doc.us_registro.username if doc.us_registro else '-',
                    'tamanho_kb': tamanho_kb,
                })

            # 3. Gera manifesto / índice textual
            manifesto_linhas = [
                "=" * 80,
                "CONVIVER - DOSSIÊ INTEGRAL DE PRONTUÁRIO MÉDICO E DOCUMENTOS LEGAIS",
                "=" * 80,
                f"PACIENTE: {paciente.nome}",
                f"CPF: {getattr(paciente, 'cpf', '-') or '-'}",
                f"DATA DE NASCIMENTO: {paciente.dt_nascimento.strftime('%d/%m/%Y') if getattr(paciente, 'dt_nascimento', None) else '-'}",
                f"DATA DA GERAÇÃO DO DOSSIÊ: {data_formatada_extenso}",
                f"SOLICITADO / GERADO POR: {user.get_full_name() or user.username if user else 'Sistema'}",
                f"TOTAL DE DOCUMENTOS ATIVOS INCLUÍDOS: {len(arquivos_do_paciente)}",
                "-" * 80,
                "ÍNDICE DE ARQUIVOS INCLUÍDOS NO PACOTE:",
                "-" * 80,
            ]

            if arquivos_do_paciente:
                for idx, item in enumerate(arquivos_do_paciente, start=1):
                    manifesto_linhas.append(
                        f"{idx:02d}. [{item['categoria']}] {item['arquivo']}\n"
                        f"    Tipo: {item['tipo']} | Data Registro: {item['data']} | Profissional: {item['usuario']} | Tam: {item['tamanho_kb']} KB"
                    )
            else:
                manifesto_linhas.append("Nenhum arquivo PDF ativo encontrado no sistema para este paciente.")

            manifesto_linhas.extend([
                "-" * 80,
                "AVISO DE CONFIDENCIALIDADE E INTEGRIDADE:",
                "Este pacote contém informações médicas protegidas por sigilo profissional,",
                "código de ética médica e pela Lei Geral de Proteção de Dados (LGPD - Lei 13.709/2018).",
                "=" * 80,
            ])

            manifesto_conteudo = "\n".join(manifesto_linhas)
            manifesto_nome = f"{pasta_raiz_paciente}Manifesto_Dossie_{nome_paciente_limpo}.txt"
            zip_file.writestr(manifesto_nome, manifesto_conteudo.encode('utf-8'))

    if total_geral_arquivos == 0:
        return None, None, 0

    buffer.seek(0)
    if multiplos_pacientes:
        nome_arquivo_zip = f"Dossie_Prontuarios_Lote_{timestamp_str}.zip"
    else:
        primeiro_nome = sanitizar_nome(list(pessoas_dict.values())[0].nome)
        nome_arquivo_zip = f"Dossie_Prontuario_{primeiro_nome}_{timestamp_str}.zip"

    return buffer.getvalue(), nome_arquivo_zip, total_geral_arquivos


def gerar_zip_relatorios_selecionados(relatorios_queryset, user=None):
    """
    Gera um arquivo ZIP contendo exatamente os relatórios selecionados na listagem,
    com nomes legíveis e sem duplicidade de caminhos.
    """
    buffer = io.BytesIO()
    total_arquivos = 0
    nomes_usados = set()
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')

    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as zip_file:
        for rel in relatorios_queryset.select_related('pessoa', 'atendimento__pessoa'):
            if not rel.relatorio:
                continue
            pdf_bytes = obter_bytes_pdf(rel.relatorio)
            if not pdf_bytes:
                continue

            paciente_nome = 'Paciente'
            if rel.pessoa:
                paciente_nome = sanitizar_nome(rel.pessoa.nome)
            elif rel.atendimento and rel.atendimento.pessoa:
                paciente_nome = sanitizar_nome(rel.atendimento.pessoa.nome)

            tipo_nome = TIPO_LABELS.get((rel.tipo or '').lower(), rel.tipo or 'relatorio')
            data_reg = rel.dt_registro.strftime('%Y%m%d') if rel.dt_registro else 'semdata'
            nome_base = f"{sanitizar_nome(tipo_nome)}_{paciente_nome}_{data_reg}_ID{rel.id}.pdf"

            contador = 1
            nome_final = nome_base
            while nome_final in nomes_usados:
                nome_final = f"{nome_base.rsplit('.', 1)[0]}_{contador}.pdf"
                contador += 1

            nomes_usados.add(nome_final)
            zip_file.writestr(nome_final, pdf_bytes)
            total_arquivos += 1

    if total_arquivos == 0:
        return None, None, 0

    buffer.seek(0)
    nome_zip = f"Relatorios_Selecionados_{timestamp_str}.zip"
    return buffer.getvalue(), nome_zip, total_arquivos


def unificar_pdfs_selecionados(relatorios_queryset, user=None):
    """
    Unifica todos os PDFs válidos selecionados em um único arquivo PDF contínuo (Merge),
    usando PyPDF2.PdfMerger.
    """
    merger = PyPDF2.PdfMerger()
    total_unificados = 0
    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')

    for rel in relatorios_queryset.select_related('pessoa', 'atendimento__pessoa'):
        if not rel.relatorio:
            continue
        pdf_bytes = obter_bytes_pdf(rel.relatorio)
        if not pdf_bytes:
            continue

        try:
            pdf_stream = io.BytesIO(pdf_bytes)
            merger.append(pdf_stream)
            total_unificados += 1
        except Exception as e:
            print(f"[ExportServices] Erro ao anexar PDF no merger: {e}")

    if total_unificados == 0:
        merger.close()
        return None, None, 0

    output_stream = io.BytesIO()
    merger.write(output_stream)
    merger.close()
    output_stream.seek(0)

    nome_pdf = f"Prontuario_Unificado_{timestamp_str}.pdf"
    return output_stream.getvalue(), nome_pdf, total_unificados


def exportar_relatorios_xlsx(relatorios_queryset, user=None):
    """
    Exporta a listagem de relatórios selecionados para uma planilha Excel (.xlsx)
    com formatação institucional, cabeçalho estilizado, bordas e larguras auto-ajustadas.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatórios Gerados"

    # Estilos
    cor_header = "1B3A57"  # Azul escuro corporativo
    fill_header = PatternFill(start_color=cor_header, end_color=cor_header, fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=10, bold=True)
    font_normal = Font(name="Calibri", size=10)
    font_title = Font(name="Calibri", size=14, bold=True, color="1B3A57")

    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    # Linha 1: Título institucional
    ws.merge_cells('A1:J1')
    cell_title = ws['A1']
    cell_title.value = "CONVIVER - RELATÓRIO ANALÍTICO DE DOCUMENTOS E PRONTUÁRIOS"
    cell_title.font = font_title
    cell_title.alignment = Alignment(vertical="center", horizontal="left")
    ws.row_dimensions[1].height = 28

    # Linha 2: Metadados da extração
    ws.merge_cells('A2:J2')
    cell_meta = ws['A2']
    data_extracao = timezone.now().strftime('%d/%m/%Y às %H:%M:%S')
    usuario_str = user.get_full_name() or user.username if user else 'Sistema'
    cell_meta.value = f"Extraído em: {data_extracao} | Solicitado por: {usuario_str}"
    cell_meta.font = Font(name="Calibri", size=9, italic=True, color="555555")
    ws.row_dimensions[2].height = 18

    # Linha 4: Cabeçalhos da Tabela
    headers = [
        "ID",
        "Tipo de Módulo",
        "Paciente",
        "CPF Paciente",
        "ID Atendimento",
        "Arquivo PDF",
        "Data de Registro",
        "Cadastrado Por",
        "Status",
        "Tamanho Aprox.",
    ]

    header_row_num = 4
    ws.row_dimensions[header_row_num].height = 24

    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_num, column=col_idx, value=header_text)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 4, 5, 7, 9, 10) else "left", vertical="center")
        cell.border = border_thin

    # Linhas de Dados
    current_row = 5
    for rel in relatorios_queryset.select_related('pessoa', 'atendimento', 'atendimento__pessoa', 'us_registro'):
        ws.row_dimensions[current_row].height = 20

        # Identificação de Paciente
        paciente_nome = "-"
        paciente_cpf = "-"
        if rel.pessoa:
            paciente_nome = rel.pessoa.nome
            paciente_cpf = getattr(rel.pessoa, 'cpf', '-') or '-'
        elif rel.atendimento and rel.atendimento.pessoa:
            paciente_nome = rel.atendimento.pessoa.nome
            paciente_cpf = getattr(rel.atendimento.pessoa, 'cpf', '-') or '-'

        # Tipo formatado
        tipo_str = TIPO_LABELS.get((rel.tipo or '').lower(), rel.tipo or 'Outro')

        # Nome do arquivo e tamanho
        nome_arquivo = rel.relatorio.name.rsplit('/', 1)[-1] if rel.relatorio else '-'
        tamanho_str = "-"
        if rel.relatorio and default_storage.exists(rel.relatorio.name):
            try:
                tam_bytes = default_storage.size(rel.relatorio.name)
                tamanho_str = f"{round(tam_bytes / 1024, 1)} KB"
            except Exception:
                pass

        # Data formatada
        data_str = rel.dt_registro.strftime('%d/%m/%Y %H:%M') if rel.dt_registro else '-'
        usuario_reg = rel.us_registro.username if rel.us_registro else '-'
        status_str = "Ativo" if rel.status == 'A' else "Inativo"
        atend_id = rel.atendimento_id if rel.atendimento_id else '-'

        row_values = [
            rel.id,
            tipo_str,
            paciente_nome,
            paciente_cpf,
            atend_id,
            nome_arquivo,
            data_str,
            usuario_reg,
            status_str,
            tamanho_str,
        ]

        is_even = (current_row % 2 == 0)
        for col_idx, val in enumerate(row_values, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = font_normal
            c.border = border_thin
            if is_even:
                c.fill = fill_zebra
            
            # Alinhamentos específicos
            if col_idx in (1, 4, 5, 7, 9, 10):
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # Auto-ajuste de largura de colunas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 4:  # pula as linhas de título/cabeçalho mesclado
                continue
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    timestamp_str = timezone.now().strftime('%Y%m%d_%H%M%S')
    nome_xlsx = f"Relatorios_Gerados_{timestamp_str}.xlsx"
    return buffer.getvalue(), nome_xlsx
